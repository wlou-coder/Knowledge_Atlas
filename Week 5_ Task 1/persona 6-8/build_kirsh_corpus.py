import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── palette ───────────────────────────────────────────────────────────────────
CHARCOAL      = "1C2833"   # header bg
SLATE_MID     = "2E4057"   # sub-flavour divider
STEEL_LIGHT   = "EAF0F6"   # alt row A
CREAM         = "FDFBF0"   # alt row B
WHITE         = "FFFFFF"
AMBER         = "FFF3CD"   # alt row B fallback

def hdr_fill(hex_color): return PatternFill("solid", fgColor=hex_color)
def thin_border():
    s = Side(style="thin", color="BBBBBB")
    return Border(left=s, right=s, top=s, bottom=s)

HEADERS = [
    "id", "question", "adequacy_condition",
    "cognitive_purpose", "answer_shape", "evidential_demand",
    "persona_fit", "theoretical_commitment",
    "source", "provenance", "notes"
]

rows = [
# ──────────────────────────────────────────────────────────────────────────────
# SUB-FLAVOUR 1 · P8-triage-operator  (D-Q-001 – D-Q-010)
# ──────────────────────────────────────────────────────────────────────────────
("D-Q-001",
 "What is the highest-priority action item for me in the next 30 minutes — is anything in the pipeline broken, blocked, or actively drifting?",
 "Fails if it does not sort items by urgency and does not distinguish between system failures, student blockages, and epistemic drift.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes triage is possible only when the system's state is visible and sortable by urgency; does not engage with items whose priority is contested.",
 "panel", "Panel-A · scenario-card-morning-triage",
 "First question on arrival; the dashboard must answer this before anything else."),

("D-Q-002",
 "Are there HITL review items that have been in the queue for more than 48 hours without action — and who is responsible for each?",
 "Fails if it does not give item age, item type, and responsible party for every overdue review.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes HITL latency is a reliable proxy for pipeline health; does not engage with legitimate reasons for delay (item complexity, student illness).",
 "panel", "Panel-A · scenario-card-HITL-queue-audit",
 ""),

("D-Q-003",
 "Did any extractor fail silently in the past 24 hours — i.e., return outputs without flagging errors I would not catch from the result alone?",
 "Fails if it does not specify extractor name, failure mode type (silent vs. flagged), and count of affected outputs.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes silent failures are the most dangerous class; flagged errors are self-reporting and lower priority.",
 "panel", "Panel-A · scenario-card-silent-failure-check",
 "Silent failure is epistemically worse than a loud crash — it corrupts the knowledge base without triggering a review."),

("D-Q-004",
 "Which students across all four tracks have not submitted any work in the past 72 hours — and for which of those is silence a concern vs. expected?",
 "Fails if it does not distinguish expected silence (no deadline, weekend) from concerning silence (deadline passed, no communication).",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes activity timestamps reliably proxy progress; does not engage with offline work or asynchronous contributions.",
 "panel", "Panel-B · scenario-card-student-activity-scan",
 ""),

("D-Q-005",
 "Has the Tagging_Contractor registry produced new inconsistency flags in the past 24 hours that I have not yet seen?",
 "Fails if it does not specify flag type (duplicate tag, contradictory assignment, undefined tag used) and the affected template or article.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes the registry's self-consistency checks are reliable; does not engage with meta-level inconsistencies the checks themselves might miss.",
 "panel", "Panel-A · scenario-card-registry-flags",
 ""),

("D-Q-006",
 "Are there warrants in the Atlas that changed state — from 'supported' to 'contested' or vice versa — in the past 24 hours as a result of new extractor runs?",
 "Fails if it does not name the specific warrant(s), the direction of state change, and the trigger (new extraction, student edit, or HITL decision).",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes warrant state is discrete (supported / contested / underdetermined) and tracked in a timestamped change log.",
 "panel", "Panel-A · scenario-card-warrant-state-changes",
 ""),

("D-Q-007",
 "Is any component of the multi-AI workflow currently running outside its normal latency range — and is that a bottleneck blocking downstream tasks or a background process?",
 "Fails if it does not give latency relative to baseline and does not distinguish bottleneck (blocking) from background (non-blocking).",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes AI workflow latency has a known stable baseline; does not engage with baseline drift as itself a diagnostic signal.",
 "panel", "Panel-A · scenario-card-AI-workflow-health",
 ""),

("D-Q-008",
 "What is the current queue depth for each of the four tracks — how many items are pending extraction, review, tagging, and HITL respectively?",
 "Fails if it does not break down queue by track and pipeline stage and does not compare current depth to the historical baseline.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes queue depth per stage is a reliable bottleneck indicator; does not engage with quality variation across items in the queue.",
 "panel", "Panel-A · scenario-card-queue-depth-scan",
 ""),

("D-Q-009",
 "If I can only do one thing in the next 20 minutes to prevent the worst downstream consequence, what is it?",
 "Fails if it does not name a single action, a specific consequence it prevents, and its ranking against alternatives.",
 "deliberation", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes the system can surface a single most-critical action; may not capture interaction effects between simultaneously critical items.",
 "panel", "Panel-A · scenario-card-single-action-triage",
 "The hardest triage question — requires the dashboard to synthesise across all five health dimensions into one recommendation."),

("D-Q-010",
 "Are there claims in the Atlas that have accumulated more than three contested annotations since last week — and what is driving the contestation?",
 "Fails if it does not name the claim, the contestation count, and the type of contestation (methodological, theoretical, or scope-dispute).",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-triage-operator",
 "Assumes contestation count proxies epistemic instability; does not distinguish productive scholarly debate from data-entry error.",
 "panel", "Panel-A · scenario-card-contested-claim-scan",
 ""),

# ──────────────────────────────────────────────────────────────────────────────
# SUB-FLAVOUR 2 · P8-epistemic-graph-reader  (D-Q-011 – D-Q-020)
# ──────────────────────────────────────────────────────────────────────────────
("D-Q-011",
 "Show me the current state of the ART-vs-SRT underdetermination node in the epistemic graph — has any new evidence shifted the balance since my last session?",
 "Fails if it does not show current evidence weight on each side and does not indicate whether any new PNU extraction changed the balance.",
 "information-seeking", "field-map", "converging",
 "P8-epistemic-graph-reader",
 "Assumes the epistemic graph encodes evidence weight for competing theoretical accounts in a queryable, version-tracked form.",
 "panel", "Panel-C · scenario-card-underdetermination-node-read",
 ""),

("D-Q-012",
 "Which domain currently has the highest density of underdetermination flags — biophilia, acoustics, colour, or another — and is that density increasing?",
 "Fails if it does not rank domains by underdetermination-flag density and does not distinguish genuine underdetermination from merely unreviewed claims.",
 "information-seeking", "ranked-brief", "converging",
 "P8-epistemic-graph-reader",
 "Assumes underdetermination is distributed unevenly across domains and that flag density is a meaningful epistemic health signal.",
 "panel", "Panel-C · scenario-card-underdetermination-density",
 ""),

("D-Q-013",
 "Are there warrant chains that currently depend on a single unreplicated study — and how many downstream claims are implicated in each case?",
 "Fails if it does not give the study, the dependent warrant(s), and a count of downstream claims that would be affected by retraction or failed replication.",
 "information-seeking", "field-map", "measurement-grade",
 "P8-epistemic-graph-reader",
 "Assumes single-study dependence is a structural vulnerability; graph depth of downstream implication is the risk metric.",
 "mining", "template-single-study-warrant-dependency",
 ""),

("D-Q-014",
 "Which PNU templates have the fewest backing studies — and are those the same templates students are currently mining for Track 4?",
 "Fails if it does not cross-reference template backing-study count against current Track 4 mining assignments.",
 "inquiry", "field-map", "converging",
 "P8-epistemic-graph-reader",
 "Assumes template robustness (backing-study count) is a design input for which templates to prioritise for student mining work.",
 "panel", "Panel-B · scenario-card-thin-template-audit",
 ""),

("D-Q-015",
 "Has any extractor been producing systematically biased extractions in a detectable direction over the past two weeks — e.g., consistently inflating effect-size confidence?",
 "Fails if it does not give a directional bias estimate and does not compare against a HITL gold-standard set.",
 "inquiry", "Toulmin", "measurement-grade",
 "P8-epistemic-graph-reader",
 "Assumes extractors can have systematic directional biases detectable via comparison to gold standard; random error would not show up as directional drift.",
 "panel", "Panel-A · scenario-card-extractor-bias-trend",
 ""),

("D-Q-016",
 "Are there nodes in the epistemic graph where two competing theoretical accounts each have a plausible warrant, but no HITL decision has been made to adjudicate?",
 "Fails if it does not list the competing warrants at each such node and does not flag whether the item is in the HITL adjudication queue.",
 "information-seeking", "field-map", "mechanistic",
 "P8-epistemic-graph-reader",
 "Assumes the graph can represent dual-warrant states (both accounts provisionally supported) and that these are distinct from simple underdetermination.",
 "panel", "Panel-C · scenario-card-dual-warrant-nodes",
 ""),

("D-Q-017",
 "How has the ratio of 'supported' to 'contested' warrant states changed over the past two weeks — is the evidence base getting more or less settled?",
 "Fails if it does not give a ratio at two time points and does not distinguish new contestation driven by new extractions vs. student edits.",
 "inquiry", "Toulmin", "measurement-grade",
 "P8-epistemic-graph-reader",
 "Assumes the supported/contested ratio is a meaningful indicator of the field's epistemic maturity as reflected in the Atlas.",
 "panel", "Panel-A · scenario-card-warrant-state-trend",
 ""),

("D-Q-018",
 "Are there claims tagged 'causal-with-mechanism' whose backing studies only establish correlation — i.e., a tagging inflation error that overstates epistemic confidence?",
 "Fails if it does not name the specific claim, the backing study design, and the reason it does not meet the causal-with-mechanism standard.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-epistemic-graph-reader",
 "Assumes the evidential-demand taxonomy is precisely defined and that mismatches between tag and study design are mechanically detectable.",
 "mining", "template-evidential-demand-tagging-audit",
 ""),

("D-Q-019",
 "Which branches of the IV × DV cross-product matrix are completely unpopulated — no PNU template exists for that combination — and does the absence reflect an evidence gap or a mining gap?",
 "Fails if it does not identify at least three unpopulated cells and does not say whether absence reflects genuine evidence gaps or artefacts of mining coverage.",
 "discovery", "field-map", "suggestive",
 "P8-epistemic-graph-reader",
 "Assumes the IV × DV matrix should in principle be exhaustively mappable; absence may be evidential or artefactual — the distinction matters for next steps.",
 "cross-product", "IV:* × DV:* · unpopulated-cells scan across all four corners",
 ""),

("D-Q-020",
 "Show me all warrants whose qualifier specifies a population restriction (e.g., 'healthy adults only') that would make the claim inapplicable to clinical populations — flagging those deployed in healthcare design contexts.",
 "Fails if it does not cross-reference qualifier population restrictions against the space-type application contexts where the claim currently appears in the Atlas.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-epistemic-graph-reader",
 "Assumes qualifier population restrictions are recorded in structured form enabling cross-referencing with application context; mismatch is an epistemic validity risk.",
 "panel", "Panel-C · scenario-card-qualifier-mismatch-audit",
 ""),

# ──────────────────────────────────────────────────────────────────────────────
# SUB-FLAVOUR 3 · P8-student-pipeline-monitor  (D-Q-021 – D-Q-030)
# ──────────────────────────────────────────────────────────────────────────────
("D-Q-021",
 "Which Track 4 student is furthest behind their expected milestone for this week — and do I need to intervene, or is there a self-correcting mechanism in place?",
 "Fails if it does not give a specific student, their current milestone status, and a recommendation on whether direct PI intervention is warranted.",
 "deliberation", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes milestone completion rate is the primary progress indicator; does not engage with quality of completed milestones vs. raw completion count.",
 "panel", "Panel-B · scenario-card-student-milestone-audit",
 ""),

("D-Q-022",
 "Across all four tracks, are there student contributions submitted but not reviewed in more than 48 hours — and is the bottleneck at HITL, TA review, or automated extraction?",
 "Fails if it does not identify the pipeline stage where each stalled contribution is sitting and does not name the responsible reviewer at that stage.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes pipeline stage is attributable to a specific responsible party; does not engage with ambiguous handoff boundaries.",
 "panel", "Panel-B · scenario-card-review-bottleneck-scan",
 ""),

("D-Q-023",
 "Is any student currently working on a template or article already processed — i.e., is there a duplication-of-effort problem I should flag before more work is wasted?",
 "Fails if it does not name the student, the overlapping item, and the other student or prior run that already processed it.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes the assignment registry is sufficiently granular to detect overlap; does not engage with intentional parallel processing as a quality check.",
 "panel", "Panel-B · scenario-card-duplication-of-effort",
 ""),

("D-Q-024",
 "Which Track 1 students have produced tagging outputs that consistently diverge from the HITL gold standard — and is there a systematic error pattern (always missing qualifiers, always over-asserting causation)?",
 "Fails if it does not give a per-student divergence pattern and does not distinguish systematic error from random noise across their submissions.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes systematic tagging errors have a detectable pattern and that the pattern is diagnostic of a specific conceptual misunderstanding.",
 "panel", "Panel-B · scenario-card-Track1-tagging-divergence",
 ""),

("D-Q-025",
 "Has any student submitted a Task 1 question corpus where adequacy conditions are formulaic — all reading as 'I will know it when I see it' — that should have been gated by the Codebook Custodian?",
 "Fails if it does not flag the specific adequacy conditions that fail the one-sentence criterion and does not note whether the Codebook Custodian reviewed them.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes the adequacy-condition quality check is operationalisable via pattern matching on vague phrasing; does not engage with intentionally open adequacy conditions.",
 "panel", "Panel-B · scenario-card-Task1-adequacy-quality-check",
 ""),

("D-Q-026",
 "Are all four Track 4 students working on genuinely distinct personas — or has any pair converged on questions too similar to provide the cross-persona contrast the assignment requires?",
 "Fails if it does not compare question corpora across pairs and does not provide a divergence metric or qualitative diagnosis of overlap.",
 "inquiry", "contrast-pair", "suggestive",
 "P8-student-pipeline-monitor",
 "Assumes cross-persona distinctiveness is measurable and pedagogically critical; does not engage with productive convergence as a signal of shared domain structure.",
 "panel", "Panel-B · scenario-card-cross-persona-divergence-check",
 ""),

("D-Q-027",
 "Which students have not yet engaged with the researcher corpus as a calibration standard — and what is the shortest intervention that would get them there?",
 "Fails if it does not identify non-engagement evidence (no citation to researcher corpus in submitted work) and does not name a concrete, minimal intervention.",
 "deliberation", "ranked-brief", "suggestive",
 "P8-student-pipeline-monitor",
 "Assumes calibration against the researcher corpus is necessary for Task 1 quality; does not engage with students who may have self-calibrated independently.",
 "panel", "Panel-B · scenario-card-researcher-corpus-calibration",
 ""),

("D-Q-028",
 "For Track 2 students, are any question-to-answer-shape fittings internally inconsistent — e.g., a question tagged 'deliberation' fitted to a Toulmin shape without justification?",
 "Fails if it does not cite the schema-fitting decision tree and does not name specific rows where the cognitive_purpose × answer_shape mismatch occurs.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes the decision tree produces a principled mapping from cognitive_purpose × question type to answer shape; inconsistency is diagnostic of conceptual confusion.",
 "mining", "template-answer-shape-consistency-audit",
 ""),

("D-Q-029",
 "Has any student's question corpus grown beyond 60 questions without a winnowing log — suggesting the four-stage winnowing protocol was skipped?",
 "Fails if it does not compare corpus size to the submitted or missing winnowing log and does not flag the protocol stage where the student appears to have stopped.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes corpus size > 60 without winnowing is a process-compliance failure, not a deliberate content decision.",
 "panel", "Panel-B · scenario-card-corpus-size-audit",
 ""),

("D-Q-030",
 "Which students are on track to deliver Task 3 prototypes with a working Chinn-Brewer panel, and which are at risk of producing a tutorial page that hides its defeaters?",
 "Fails if it does not give a per-student assessment and does not name the specific evidence (or absence) for the rebuttal-panel element in each current prototype sketch.",
 "deliberation", "ranked-brief", "measurement-grade",
 "P8-student-pipeline-monitor",
 "Assumes presence of a Chinn-Brewer panel is the critical design move separating an evidential journey from a tutorial; absence is a design failure, not an optional omission.",
 "panel", "Panel-B · scenario-card-Task3-rebuttal-risk",
 ""),

# ──────────────────────────────────────────────────────────────────────────────
# SUB-FLAVOUR 4 · P8-quality-controller  (D-Q-031 – D-Q-040)
# ──────────────────────────────────────────────────────────────────────────────
("D-Q-031",
 "What is the current inter-rater reliability between the automated extractor and the HITL gold standard, broken down by extraction field?",
 "Fails if it does not give a per-field reliability metric (Cohen's κ or equivalent) and does not compare to the last recorded baseline.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes IRR is the right metric for extractor quality; does not engage with the possibility that the HITL gold standard itself has a systematic bias.",
 "panel", "Panel-A · scenario-card-extractor-IRR",
 ""),

("D-Q-032",
 "Has the Tagging_Contractor's tag-assignment accuracy drifted from the week-3 baseline — and in which domains is the drift largest?",
 "Fails if it does not specify the drift metric (accuracy delta), the direction (over- or under-tagging), and the domain with the largest drift.",
 "inquiry", "Toulmin", "measurement-grade",
 "P8-quality-controller",
 "Assumes tagging accuracy has a stable baseline against which drift is meaningful; does not engage with legitimate evolution of the tagging scheme as a drift cause.",
 "panel", "Panel-A · scenario-card-tagging-drift",
 ""),

("D-Q-033",
 "Are there PNU template extractions where the extractor's confidence score is higher than the HITL reviewer's subsequent assessment — indicating systematic overconfidence?",
 "Fails if it does not give a count of overconfident extractions, the magnitude of the gap, and whether overconfidence concentrates in a particular extractor or article type.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes extractor confidence scores are calibrated and that systematic overconfidence is diagnosable via HITL comparison; does not engage with cases where HITL is wrong.",
 "panel", "Panel-A · scenario-card-extractor-overconfidence",
 ""),

("D-Q-034",
 "Which tags in the Tagging_Contractor registry are currently being used inconsistently — applied to articles that do not match their definition according to HITL review?",
 "Fails if it does not name the specific tags, the articles where misapplication occurred, and the frequency of misapplication.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes tag definitions are stable and misapplication is diagnosable via comparison to definition; does not engage with definition ambiguity as a cause.",
 "mining", "template-tag-misapplication-audit",
 ""),

("D-Q-035",
 "Are there HITL review decisions in the past week that were later reversed — and what type of claim reversal is most common (effect direction, effect size, scope restriction)?",
 "Fails if it does not categorise reversals by type and does not note whether any reversal affected a downstream claim's warrant state.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes HITL reversals are trackable events with a claim type and a downstream consequence; does not engage with reversals that cancel out.",
 "panel", "Panel-A · scenario-card-HITL-reversal-audit",
 ""),

("D-Q-036",
 "Has the Article_Eater extractor's precision on effect-size extraction improved, stayed flat, or declined over the past three weeks?",
 "Fails if it does not give a three-week trend line (not just a current snapshot) and does not identify any confounding variable (article type, field, length).",
 "inquiry", "Toulmin", "measurement-grade",
 "P8-quality-controller",
 "Assumes extractor precision has a meaningful trend interpretable without knowing article-type distribution changes over the same period.",
 "panel", "Panel-A · scenario-card-Article-Eater-trend",
 ""),

("D-Q-037",
 "Are there claims where the warrant tag was assigned by a student and has never been HITL-reviewed — i.e., unreviewed student-assigned epistemic status in the live Atlas?",
 "Fails if it does not give a count of unreviewed student-assigned warrant tags and does not flag those in domains where the student's tagging track record is poor.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes student-assigned warrant tags carry lower reliability than HITL-reviewed tags; unreviewed tags are a latent quality risk in the live knowledge base.",
 "panel", "Panel-A · scenario-card-unreviewed-warrant-tags",
 ""),

("D-Q-038",
 "What percentage of PNU templates currently in the Atlas have at least two independent backing studies — and has that percentage changed since the term started?",
 "Fails if it does not give the percentage at two time points and does not distinguish templates with zero, one, and two-or-more backing studies separately.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes two independent backing studies is the minimum threshold for a template to be reliably cited in a student evidential journey.",
 "mining", "template-backing-study-coverage",
 ""),

("D-Q-039",
 "Are there domain ontologies in the Outcome_Contractor where term definitions have drifted — e.g., 'affect' terms now being used to code cognitive outcomes?",
 "Fails if it does not name the specific terms and articles where drift is observed and does not give a frequency count of misapplications.",
 "inquiry", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes ontological drift is detectable via comparison to original term definitions; does not engage with legitimate ontological evolution vs. error.",
 "panel", "Panel-C · scenario-card-ontological-drift",
 ""),

("D-Q-040",
 "Has any student extraction or tagging work been incorporated into the live Atlas without a HITL review step — i.e., is the quality gate being bypassed anywhere?",
 "Fails if it does not name the specific contribution, the pipeline step where review was skipped, and the commit or timestamp.",
 "information-seeking", "ranked-brief", "measurement-grade",
 "P8-quality-controller",
 "Assumes the HITL gate is both necessary and enforceable; does not engage with deliberate bypass decisions made under time pressure.",
 "panel", "Panel-A · scenario-card-quality-gate-bypass",
 ""),

# ──────────────────────────────────────────────────────────────────────────────
# SUB-FLAVOUR 5 · P8-site-architect  (D-Q-041 – D-Q-050)
# ──────────────────────────────────────────────────────────────────────────────
("D-Q-041",
 "What should appear in the first 'above the fold' zone of ka_pi_dashboard.html to give me an accurate 90-second triage read — and what should be deferred to a second scroll?",
 "Fails if it does not name specific display elements with priority ordering and does not distinguish triage-required information from deeper-investigation information.",
 "deliberation", "procedure", "suggestive",
 "P8-site-architect",
 "Assumes a 90-second constraint implies a specific hierarchy of information density; draws on attention-economy and information-scent design principles.",
 "panel", "Panel-C · scenario-card-dashboard-above-fold",
 "This is the core deliverable question — ka_pi_dashboard.html does not exist yet; designing it is one of the quarter's concrete outputs."),

("D-Q-042",
 "How should the HITL adjudication queue be sorted and displayed on the PI dashboard — by item age, by claim type, by downstream-consequence count, or by student?",
 "Fails if it does not give a primary and secondary sort key and does not justify each with a consequence argument.",
 "deliberation", "contrast-pair", "suggestive",
 "P8-site-architect",
 "Assumes queue sorting is a design decision with epistemic consequences — different sort orders surface different risks first.",
 "panel", "Panel-C · scenario-card-HITL-queue-display",
 ""),

("D-Q-043",
 "What alert threshold should trigger an automatic notification for extractor reliability degradation — and how do I avoid alert fatigue from over-sensitive thresholds?",
 "Fails if it does not name a specific metric, a threshold value, and a false-positive rate estimate for that threshold.",
 "deliberation", "contrast-pair", "measurement-grade",
 "P8-site-architect",
 "Assumes a single fixed threshold is the right alert design; does not engage with adaptive or context-sensitive thresholds.",
 "panel", "Panel-C · scenario-card-alert-threshold-design",
 ""),

("D-Q-044",
 "Should the PI dashboard display individual student names and work states, or aggregate track-level statistics — and what are the pedagogical and privacy trade-offs?",
 "Fails if it does not weigh the pedagogical benefit of individual visibility against the privacy cost and does not give a design recommendation.",
 "deliberation", "contrast-pair", "suggestive",
 "P8-site-architect",
 "Assumes individual visibility improves triage accuracy but creates a surveillance dynamic; does not engage with hybrid designs that expose aggregate by default.",
 "panel", "Panel-C · scenario-card-dashboard-privacy-design",
 ""),

("D-Q-045",
 "How should the dashboard represent underdetermination density across the epistemic graph — a heat map, sorted list, graph visualisation, or plain-text alert?",
 "Fails if it does not give a primary recommendation justified by the 90-second readability constraint and Kirsh's triage goal.",
 "deliberation", "ranked-brief", "suggestive",
 "P8-site-architect",
 "Assumes triage efficiency requires a specific visual encoding; different representations afford different inspection strategies under time pressure.",
 "panel", "Panel-C · scenario-card-underdetermination-visualisation",
 ""),

("D-Q-046",
 "What should the 'warrant drift' indicator on the dashboard actually measure — state changes per day, contested-to-supported ratio, or another operationalisation?",
 "Fails if it does not give a specific operationalisation and does not explain why it is preferable to alternatives for the PI's triage purpose.",
 "deliberation", "contrast-pair", "measurement-grade",
 "P8-site-architect",
 "Assumes warrant drift is a meaningful health signal and that the operationalisation choice encodes assumptions about what 'healthy' drift looks like.",
 "cross-product", "IV:system[warrant-state-change] × DV:system[triage-accuracy] · dashboard-design corner",
 ""),

("D-Q-047",
 "How should the dashboard handle multiple simultaneously urgent items competing for my attention — force a single recommended action or present a ranked list?",
 "Fails if it does not address the cognitive cost of simultaneous urgent items and does not give a recommendation grounded in decision-support or cognitive-load research.",
 "deliberation", "contrast-pair", "suggestive",
 "P8-site-architect",
 "Assumes forced-single-action and ranked-list designs have different effects on decision quality under cognitive load; tests which design principle to apply.",
 "panel", "Panel-C · scenario-card-multi-urgent-display",
 ""),

("D-Q-048",
 "What does the 'things the system wants me to adjudicate' queue need to look like for me to trust it — what metadata per item would make me confident it is triaged correctly before I see it?",
 "Fails if it does not specify at least four metadata fields (age, claim type, downstream count, triggering event) and does not address the trust question explicitly.",
 "deliberation", "procedure", "suggestive",
 "P8-site-architect",
 "Assumes trust in an automated triage queue requires specific metadata; absence of metadata induces distrust regardless of underlying accuracy.",
 "panel", "Panel-C · scenario-card-adjudication-queue-trust",
 ""),

("D-Q-049",
 "How should the dashboard signal the difference between a system error (something technically broken) and an epistemic error (a claim has drifted) — and should these use the same alert affordance?",
 "Fails if it does not distinguish the two error types and does not give a design reason for using different or identical alert affordances.",
 "deliberation", "contrast-pair", "suggestive",
 "P8-site-architect",
 "Assumes system errors and epistemic errors require different response modes; conflating them would create category confusion in triage.",
 "panel", "Panel-C · scenario-card-error-type-signalling",
 ""),

("D-Q-050",
 "If the PI dashboard can only show five data points before I click through, what are those five — and how would I validate that they give an accurate health read rather than a misleading one?",
 "Fails if it does not name five specific data points, rank-order them with justification, and describe a validation test that would detect false-green or false-red readings.",
 "deliberation", "ranked-brief", "measurement-grade",
 "P8-site-architect",
 "Assumes a minimal dashboard is superior to an overwhelming one; the five-item constraint encodes a view about cognitive load in triage contexts.",
 "panel", "Panel-C · scenario-card-minimal-viable-dashboard",
 "The architectural closure question — defines what the dashboard fundamentally IS."),
]

# ── workbook build ─────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "P8-Kirsh-Corpus"

col_widths = [10, 62, 57, 22, 20, 24, 26, 57, 16, 52, 45]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 30
for col_idx, hdr in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col_idx, value=hdr)
    cell.fill      = hdr_fill(CHARCOAL)
    cell.font      = Font(bold=True, color=WHITE, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = thin_border()

subflavour_map = {
    2:  "① TRIAGE-OPERATOR — What needs my attention right now? Is anything broken?",
    13: "② EPISTEMIC-GRAPH-READER — What is the current state of the evidence graph?",
    24: "③ STUDENT-PIPELINE-MONITOR — Who is stuck? What has been submitted? What needs review?",
    35: "④ QUALITY-CONTROLLER — Are extractors reliable? Is tagging consistent?",
    46: "⑤ SITE-ARCHITECT — How should the PI dashboard and pipeline be designed?",
}

current_data_row = 2
for r_idx, row_data in enumerate(rows):
    if current_data_row in subflavour_map:
        ws.row_dimensions[current_data_row].height = 22
        label_cell = ws.cell(row=current_data_row, column=1,
                             value=subflavour_map[current_data_row])
        label_cell.fill      = hdr_fill(SLATE_MID)
        label_cell.font      = Font(bold=True, color=WHITE, size=10)
        label_cell.alignment = Alignment(vertical="center")
        label_cell.border    = thin_border()
        ws.merge_cells(start_row=current_data_row, start_column=1,
                       end_row=current_data_row, end_column=len(HEADERS))
        current_data_row += 1

    fill_color = STEEL_LIGHT if (r_idx % 2 == 0) else AMBER
    ws.row_dimensions[current_data_row].height = 90

    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=current_data_row, column=col_idx, value=value)
        cell.fill      = hdr_fill(fill_color)
        cell.border    = thin_border()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font      = Font(bold=(col_idx == 1), size=10)

    current_data_row += 1

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

out_path = "/sessions/confident-magical-goldberg/mnt/outputs/P8_David_Kirsh_Question_Corpus.xlsx"
wb.save(out_path)
print(f"Saved → {out_path}")
print(f"Total question rows: {len(rows)}")
